#!/usr/bin/env python3
"""Core chemistry and XLSX calculation engine for RRF Calculator v2.

Dictionary schema:
    name, type, formula_rule, summary_group

Examples:
    {"name": "aromatic", "type": "class", "formula_rule": "C[n]H[2*n-6]", ...}
    {"name": "benzoic acid", "type": "compound", "formula_rule": "C7H6O2", ...}

A ``class`` entry requires ``C<number>`` in Group Name and expands a formula
rule. A ``compound`` entry is an exact molecular formula and does not require a
carbon-number suffix.
"""

from __future__ import annotations

import argparse
import ast
import json
import math
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font


NORMALIZATION_FACTOR = 0.83905
ATOMIC_WEIGHTS = {
    "H": 1.008,
    "B": 10.81,
    "C": 12.011,
    "N": 14.007,
    "O": 15.999,
    "F": 18.998403163,
    "Si": 28.085,
    "P": 30.973761998,
    "S": 32.06,
    "Cl": 35.45,
    "Br": 79.904,
    "I": 126.90447,
}
C_AW = ATOMIC_WEIGHTS["C"]

SUMMARY_SHEET_NAME = "Class summary"
SUMMARY_C_MIN = 7
SUMMARY_C_MAX = 25
VALID_ENTRY_TYPES = {"class", "compound"}


class MissingCarbonNumber(Exception):
    """Raised when a class entry is used without an explicit C number."""


class UnknownChemical(Exception):
    """Raised when Group Name has no matching dictionary entry."""


GROUP_RE = re.compile(
    r"^(?P<name>.+?)(?:\s+C(?P<C>\d+)(?P<suffix>\+|\s*\?)?)?$",
    re.IGNORECASE,
)
EXACT_TOKEN_RE = re.compile(r"([A-Z][a-z]?)(\d*)")
TEMPLATE_TOKEN_RE = re.compile(r"([A-Z][a-z]?)(?:\[([^\]]+)\]|(\d*))")


@dataclass(frozen=True)
class ResolvedChemical:
    dictionary_name: str
    entry_type: str
    formula: str
    elements: dict[str, int]
    molecular_weight: float
    carbon_count: int
    hydrogen_count: int
    suffix: str | None
    summary_group: str


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def _safe_numeric_eval(expression: str, variable_name: str, variable_value: int) -> float:
    """Evaluate a small arithmetic expression without eval()."""
    tree = ast.parse(expression, mode="eval")

    def walk(node):
        if isinstance(node, ast.Expression):
            return walk(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name) and node.id == variable_name:
            return variable_value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = walk(node.operand)
            return value if isinstance(node.op, ast.UAdd) else -value
        if isinstance(node, ast.BinOp):
            left = walk(node.left)
            right = walk(node.right)
            if isinstance(node.op, ast.Add):
                return left + right
            if isinstance(node.op, ast.Sub):
                return left - right
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Div):
                return left / right
            if isinstance(node.op, ast.Pow):
                return left ** right
        raise ValueError(f"Unsupported expression: {expression!r}")

    return float(walk(tree))


def safe_formula_eval(expression: str, C: int) -> float:
    """Backward-compatible helper for old H_formula expressions."""
    return _safe_numeric_eval(expression, "C", C)


def safe_template_eval(expression: str, n: int) -> float:
    return _safe_numeric_eval(expression, "n", n)


def _integer_count(value: float, context: str) -> int:
    if not math.isfinite(value) or not math.isclose(value, round(value), abs_tol=1e-9):
        raise ValueError(f"Element count must be an integer in {context!r}: {value}")
    count = int(round(value))
    if count < 0:
        raise ValueError(f"Element count cannot be negative in {context!r}: {count}")
    return count


def parse_molecular_formula(formula: str) -> tuple[dict[str, int], list[str]]:
    """Parse a simple molecular formula such as C7H6O2 or C6H5OH.

    Repeated element tokens are summed. Parentheses, hydrates and charges are
    deliberately not accepted so invalid dictionary entries fail loudly.
    """
    text = formula.strip().replace(" ", "")
    if not text:
        raise ValueError("Molecular formula is empty.")

    counts: dict[str, int] = {}
    order: list[str] = []
    pos = 0
    for match in EXACT_TOKEN_RE.finditer(text):
        if match.start() != pos:
            raise ValueError(f"Unsupported molecular formula syntax: {formula!r}")
        element, number = match.groups()
        if element not in ATOMIC_WEIGHTS:
            raise ValueError(f"Unsupported element {element!r} in formula {formula!r}")
        count = int(number) if number else 1
        if element not in counts:
            order.append(element)
            counts[element] = 0
        counts[element] += count
        pos = match.end()

    if pos != len(text):
        raise ValueError(f"Unsupported molecular formula syntax: {formula!r}")
    return counts, order


def format_molecular_formula(elements: dict[str, int], order: list[str] | None = None) -> str:
    if order is None:
        # Hill-system-like fallback: C, H, then the rest alphabetically.
        order = []
        if "C" in elements:
            order.append("C")
        if "H" in elements:
            order.append("H")
        order.extend(sorted(k for k in elements if k not in {"C", "H"}))

    pieces = []
    for element in order:
        count = int(elements.get(element, 0))
        if count <= 0:
            continue
        pieces.append(element if count == 1 else f"{element}{count}")
    if not pieces:
        raise ValueError("Molecular formula contains no atoms.")
    return "".join(pieces)


def instantiate_formula_rule(rule: str, n: int) -> tuple[str, dict[str, int]]:
    """Expand a class rule like C[n]H[2*n-6]O to a concrete formula."""
    text = rule.strip().replace(" ", "")
    if not text:
        raise ValueError("Formula rule is empty.")

    counts: dict[str, int] = {}
    order: list[str] = []
    pos = 0

    for match in TEMPLATE_TOKEN_RE.finditer(text):
        if match.start() != pos:
            raise ValueError(f"Unsupported formula rule syntax: {rule!r}")
        element, expression, literal_count = match.groups()
        if element not in ATOMIC_WEIGHTS:
            raise ValueError(f"Unsupported element {element!r} in formula rule {rule!r}")

        if expression is not None:
            count = _integer_count(safe_template_eval(expression, n), rule)
        elif literal_count:
            count = int(literal_count)
        else:
            count = 1

        if element not in counts:
            order.append(element)
            counts[element] = 0
        counts[element] += count
        pos = match.end()

    if pos != len(text):
        raise ValueError(f"Unsupported formula rule syntax: {rule!r}")

    formula = format_molecular_formula(counts, order)
    return formula, counts


def molecular_weight(elements: dict[str, int]) -> float:
    total = 0.0
    for element, count in elements.items():
        try:
            atomic_weight = ATOMIC_WEIGHTS[element]
        except KeyError as exc:
            raise ValueError(f"Unsupported element {element!r}") from exc
        total += atomic_weight * count
    return total


def calculate_rrf(mw: float, carbon_count: int) -> float:
    """Calculate RRF using the existing model generalized to full MW.

    For hydrocarbons this is exactly equivalent to the previous implementation:
        ((12.011*C + 1.008*H) * 0.83905) / (12.011*C)

    For heteroatom-containing compounds the full molecular weight is used in
    the numerator. Keep this function isolated so the analytical method can be
    changed later without touching parsing/UI code.
    """
    if carbon_count <= 0:
        raise ValueError("RRF requires at least one carbon atom.")
    return mw * NORMALIZATION_FACTOR / (C_AW * carbon_count)


def migrate_legacy_dictionary(rows: list[dict]) -> list[dict[str, str]]:
    """Convert old class/H_formula dictionaries to the v2 schema."""
    migrated: list[dict[str, str]] = []
    for row in rows:
        if {"name", "type", "formula_rule", "summary_group"}.issubset(row):
            migrated.append(
                {
                    "name": normalize_text(row.get("name", "")),
                    "type": normalize_text(row.get("type", "")).casefold(),
                    "formula_rule": normalize_text(row.get("formula_rule", "")),
                    "summary_group": normalize_text(row.get("summary_group", "")),
                }
            )
            continue

        if "class" in row and "H_formula" in row:
            class_name = normalize_text(row.get("class", ""))
            h_formula = normalize_text(row.get("H_formula", ""))
            # Convert C in the old arithmetic rule to n in the new template.
            h_expression = re.sub(r"\bC\b", "n", h_formula)
            migrated.append(
                {
                    "name": class_name,
                    "type": "class",
                    "formula_rule": f"C[n]H[{h_expression}]",
                    "summary_group": normalize_text(row.get("summary_group", "")),
                }
            )
            continue

        raise ValueError("Dictionary contains an unsupported row format.")
    return migrated


def validate_dictionary(rows: list[dict]) -> list[dict[str, str]]:
    normalized = migrate_legacy_dictionary(rows)
    seen: set[tuple[str, str]] = set()

    for i, row in enumerate(normalized, start=1):
        name = normalize_text(row.get("name", ""))
        entry_type = normalize_text(row.get("type", "")).casefold()
        formula_rule = normalize_text(row.get("formula_rule", ""))
        summary_group = normalize_text(row.get("summary_group", ""))

        if not name or not entry_type or not formula_rule or not summary_group:
            raise ValueError(
                f"Dictionary row {i}: name, type, formula_rule and summary_group must be filled."
            )
        if entry_type not in VALID_ENTRY_TYPES:
            raise ValueError(
                f"Dictionary row {i}: type must be 'class' or 'compound', got {entry_type!r}."
            )

        key = (name.casefold(), entry_type)
        if key in seen:
            raise ValueError(f"Duplicate {entry_type} entry in dictionary: {name!r}")
        seen.add(key)

        if entry_type == "class":
            formula, elements = instantiate_formula_rule(formula_rule, 10)
        else:
            if "[" in formula_rule or "]" in formula_rule:
                raise ValueError(
                    f"Dictionary row {i}: compound formulas must be exact, got {formula_rule!r}."
                )
            elements, order = parse_molecular_formula(formula_rule)
            formula = format_molecular_formula(elements, order)

        if elements.get("C", 0) <= 0:
            raise ValueError(
                f"Dictionary row {i}: {name!r} resolves to {formula!r}, which has no carbon."
            )
        molecular_weight(elements)  # validates element support

        row["name"] = name
        row["type"] = entry_type
        row["formula_rule"] = formula_rule
        row["summary_group"] = summary_group

    return normalized


def build_dictionary_index(rows: list[dict]) -> dict[str, dict[str, dict[str, str]]]:
    validated = validate_dictionary(rows)
    index: dict[str, dict[str, dict[str, str]]] = {}
    for row in validated:
        index.setdefault(row["name"].casefold(), {})[row["type"]] = row
    return index


def load_dictionary_json(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8") as f:
        rows = json.load(f)
    if not isinstance(rows, list):
        raise ValueError("Dictionary JSON must contain a list of entries.")
    return validate_dictionary(rows)


def resolve_group_name(group_name: object, dictionary_rows: list[dict]) -> ResolvedChemical:
    text = normalize_text(group_name)
    match = GROUP_RE.fullmatch(text)
    if not match:
        raise ValueError(f"Cannot parse Group Name {text!r}")

    base_name = normalize_text(match.group("name"))
    key = base_name.casefold()
    index = build_dictionary_index(dictionary_rows)
    variants = index.get(key)
    if not variants:
        raise UnknownChemical(
            f"Unknown chemical/class {base_name!r}. Add it to Chemical dictionary."
        )

    c_text = match.group("C")
    suffix = match.group("suffix")

    if c_text:
        entry = variants.get("class")
        if entry is None:
            if "compound" in variants:
                raise ValueError(
                    f"{base_name!r} is an exact compound and must be used without C<number>."
                )
            raise UnknownChemical(
                f"No class rule for {base_name!r}. Add a class entry to Chemical dictionary."
            )
        n = int(c_text)
        formula, elements = instantiate_formula_rule(entry["formula_rule"], n)
    else:
        entry = variants.get("compound")
        if entry is not None:
            elements, order = parse_molecular_formula(entry["formula_rule"])
            formula = format_molecular_formula(elements, order)
        elif "class" in variants:
            raise MissingCarbonNumber(
                f"Group Name {text!r} is a class and has no explicit C number."
            )
        else:
            raise UnknownChemical(
                f"Unknown chemical {base_name!r}. Add it to Chemical dictionary."
            )

    mw = molecular_weight(elements)
    carbon_count = int(elements.get("C", 0))
    hydrogen_count = int(elements.get("H", 0))
    if carbon_count <= 0:
        raise ValueError(f"Formula {formula!r} contains no carbon; RRF cannot be calculated.")

    return ResolvedChemical(
        dictionary_name=entry["name"],
        entry_type=entry["type"],
        formula=formula,
        elements=dict(elements),
        molecular_weight=mw,
        carbon_count=carbon_count,
        hydrogen_count=hydrogen_count,
        suffix=suffix,
        summary_group=entry["summary_group"],
    )


def find_header_row_and_columns(ws):
    required = {"group name", "volume"}
    for row_idx in range(1, min(ws.max_row, 20) + 1):
        headers = {}
        for cell in ws[row_idx]:
            if cell.value is not None:
                headers[normalize_text(cell.value).casefold()] = cell.column
        if required.issubset(headers):
            return row_idx, headers
    raise ValueError("Could not find a header row containing both 'Group Name' and 'Volume'.")


def get_or_create_column(ws, header_row: int, header_name: str, style_source_col: int):
    for cell in ws[header_row]:
        if cell.value is not None and normalize_text(cell.value).casefold() == header_name.casefold():
            return cell.column

    new_col = ws.max_column + 1
    target = ws.cell(header_row, new_col)
    source = ws.cell(header_row, style_source_col)
    target.value = header_name

    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    return new_col


def create_summary_sheet(wb, calculated_rows: list[dict], total_rrf_v: float, summary_columns: list[str]) -> None:
    if SUMMARY_SHEET_NAME in wb.sheetnames:
        del wb[SUMMARY_SHEET_NAME]
    ws_summary = wb.create_sheet(SUMMARY_SHEET_NAME)

    observed_c = [int(item["C"]) for item in calculated_rows]
    c_min = min([SUMMARY_C_MIN] + observed_c) if observed_c else SUMMARY_C_MIN
    c_max = max([SUMMARY_C_MAX] + observed_c) if observed_c else SUMMARY_C_MAX
    carbon_numbers = list(range(c_min, c_max + 1))

    summary = {c: {group: 0.0 for group in summary_columns} for c in carbon_numbers}
    for item in calculated_rows:
        c = int(item["C"])
        group = str(item["summary_group"])
        summary[c][group] += item["RRF*V"] / total_rrf_v * 100.0

    headers = ["C"] + summary_columns
    for col_idx, value in enumerate(headers, start=1):
        cell = ws_summary.cell(row=1, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first_data_row = 2
    for offset, c in enumerate(carbon_numbers):
        row = first_data_row + offset
        ws_summary.cell(row=row, column=1, value=c)
        for col_idx, group in enumerate(summary_columns, start=2):
            value = summary[c][group]
            cell = ws_summary.cell(row=row, column=col_idx)
            if math.isclose(value, 0.0, abs_tol=1e-15):
                cell.value = "-"
            else:
                cell.value = value
                cell.number_format = "0.00"

    totals_row = first_data_row + len(carbon_numbers) + 1
    ws_summary.cell(row=totals_row, column=1, value="Total").font = Font(bold=True)
    grand_total = 0.0
    for col_idx, group in enumerate(summary_columns, start=2):
        total = sum(summary[c][group] for c in carbon_numbers)
        grand_total += total
        cell = ws_summary.cell(row=totals_row, column=col_idx, value=total)
        cell.number_format = "0.00"
        cell.font = Font(bold=True)

    grand_total_row = totals_row + 2
    ws_summary.cell(row=grand_total_row, column=1, value="Grand total").font = Font(bold=True)
    grand_cell = ws_summary.cell(row=grand_total_row, column=2, value=grand_total)
    grand_cell.number_format = "0.00"
    grand_cell.font = Font(bold=True)

    if summary_columns and carbon_numbers:
        first_cell = ws_summary.cell(first_data_row, 2).coordinate
        last_cell = ws_summary.cell(first_data_row + len(carbon_numbers) - 1, 1 + len(summary_columns)).coordinate
        ws_summary.conditional_formatting.add(
            f"{first_cell}:{last_cell}",
            ColorScaleRule(
                start_type="min",
                start_color="63BE7B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="F8696B",
            ),
        )

    ws_summary.freeze_panes = "B2"
    ws_summary.column_dimensions["A"].width = 14
    for col_idx in range(2, len(headers) + 1):
        ws_summary.column_dimensions[ws_summary.cell(1, col_idx).column_letter].width = 28
    ws_summary.row_dimensions[1].height = 42


def calculate_file(input_path: Path, dictionary_rows: list[dict], output_path: Path) -> dict:
    dictionary_rows = validate_dictionary(dictionary_rows)

    wb = load_workbook(input_path)
    ws = wb.active
    header_row, headers = find_header_row_and_columns(ws)
    group_col = headers["group name"]
    volume_col = headers["volume"]

    col_C = get_or_create_column(ws, header_row, "C", volume_col)
    col_H = get_or_create_column(ws, header_row, "H", volume_col)
    col_formula = get_or_create_column(ws, header_row, "Formula", volume_col)
    col_mw = get_or_create_column(ws, header_row, "Molecular Weight", volume_col)
    col_rrf = get_or_create_column(ws, header_row, "RRF", volume_col)
    col_rrfv = get_or_create_column(ws, header_row, "RRF*V", volume_col)
    col_w = get_or_create_column(ws, header_row, "w, %", volume_col)

    calculated_rows = []
    errors = []
    suffix_rows = []
    skipped_no_c = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        group_name = ws.cell(row_idx, group_col).value
        volume = ws.cell(row_idx, volume_col).value
        if group_name is None and volume is None:
            continue
        if group_name is None or normalize_text(group_name) == "":
            errors.append(f"row {row_idx}: empty Group Name")
            continue
        try:
            volume = float(volume)
        except (TypeError, ValueError):
            errors.append(f"row {row_idx}: invalid Volume = {volume!r}")
            continue

        try:
            resolved = resolve_group_name(group_name, dictionary_rows)
        except MissingCarbonNumber:
            skipped_no_c.append((row_idx, normalize_text(group_name)))
            continue
        except Exception as exc:
            errors.append(f"row {row_idx}: {exc}")
            continue

        if resolved.suffix:
            suffix_rows.append((row_idx, normalize_text(group_name)))

        rrf = calculate_rrf(resolved.molecular_weight, resolved.carbon_count)
        rrf_v = rrf * volume
        calculated_rows.append(
            {
                "row": row_idx,
                "name": resolved.dictionary_name,
                "summary_group": resolved.summary_group,
                "C": resolved.carbon_count,
                "H": resolved.hydrogen_count,
                "Formula": resolved.formula,
                "Molecular Weight": resolved.molecular_weight,
                "RRF": rrf,
                "RRF*V": rrf_v,
            }
        )

    if errors:
        preview = "\n".join(errors[:30])
        more = "" if len(errors) <= 30 else f"\n... and {len(errors) - 30} more error(s)"
        raise RuntimeError(
            f"Calculation stopped: {len(errors)} row(s) could not be processed.\n{preview}{more}"
        )

    total_rrf_v = sum(item["RRF*V"] for item in calculated_rows)
    if total_rrf_v == 0:
        raise RuntimeError("Sum(RRF*V) is zero; w, % cannot be calculated.")

    for item in calculated_rows:
        row_idx = item["row"]
        ws.cell(row_idx, col_C).value = item["C"]
        ws.cell(row_idx, col_H).value = item["H"]
        ws.cell(row_idx, col_formula).value = item["Formula"]
        ws.cell(row_idx, col_mw).value = item["Molecular Weight"]
        ws.cell(row_idx, col_rrf).value = item["RRF"]
        ws.cell(row_idx, col_rrfv).value = item["RRF*V"]
        ws.cell(row_idx, col_w).value = item["RRF*V"] / total_rrf_v * 100.0

        ws.cell(row_idx, col_C).number_format = "0"
        ws.cell(row_idx, col_H).number_format = "0"
        ws.cell(row_idx, col_mw).number_format = "0.000"
        ws.cell(row_idx, col_rrf).number_format = "0.0000"
        ws.cell(row_idx, col_rrfv).number_format = "0.00E+00"
        ws.cell(row_idx, col_w).number_format = "0.00"

    summary_columns = list(dict.fromkeys(row["summary_group"] for row in dictionary_rows))
    create_summary_sheet(wb, calculated_rows, total_rrf_v, summary_columns)
    wb.save(output_path)

    return {
        "calculated_rows": len(calculated_rows),
        "sum_rrf_v": total_rrf_v,
        "sum_w": sum(item["RRF*V"] / total_rrf_v * 100 for item in calculated_rows),
        "suffix_rows": suffix_rows,
        "skipped_no_c": skipped_no_c,
        "summary_sheet": SUMMARY_SHEET_NAME,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Calculate molecular formula, molecular weight, RRF, RRF*V, w,% and class summary."
    )
    parser.add_argument("input", type=Path, help="Input .xlsx file")
    parser.add_argument("dictionary", type=Path, help="Chemical dictionary JSON")
    parser.add_argument("-o", "--output", type=Path, default=None)
    args = parser.parse_args()

    output = args.output or args.input.with_name(args.input.stem + "_calculated.xlsx")
    rows = load_dictionary_json(args.dictionary)
    stats = calculate_file(args.input, rows, output)

    print(f"Done: {output}")
    print(f"Calculated rows: {stats['calculated_rows']}")
    print(f"Sum(RRF*V): {stats['sum_rrf_v']:.6f}")
    print(f"Sum(w, %): {stats['sum_w']:.10f}")
    print(f"Summary sheet: {stats['summary_sheet']}")
    if stats["suffix_rows"]:
        print(f"Warning: {len(stats['suffix_rows'])} row(s) contain '+' or '?'.")
    if stats["skipped_no_c"]:
        print(f"Skipped {len(stats['skipped_no_c'])} class row(s) without explicit C number.")


if __name__ == "__main__":
    main()
