from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import Qt, QStandardPaths
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QPlainTextEdit,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

import core


class RRFWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("RRF Calculator — v2")
        self.resize(1080, 720)
        self.setAcceptDrops(True)

        self.input_path: Path | None = None
        self.default_dictionary = self._load_default_dictionary()
        self.dictionary_rows = self._load_user_dictionary()

        self.tabs = QTabWidget()
        self.tabs.addTab(self._build_calculation_tab(), "Calculation")
        self.tabs.addTab(self._build_dictionary_tab(), "Chemical dictionary")
        self.setCentralWidget(self.tabs)

        self._fill_dictionary_table()
        self._log("Ready. Choose or drop an .xlsx file.")

    # -------------------- resources / persistence --------------------

    def _resource_path(self, filename: str) -> Path:
        if getattr(sys, "frozen", False):
            base = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
        else:
            base = Path(__file__).resolve().parent
        return base / filename

    def _load_default_dictionary(self) -> list[dict[str, str]]:
        with self._resource_path("default_classes.json").open("r", encoding="utf-8") as f:
            rows = json.load(f)
        return core.validate_dictionary(rows)

    def _config_dir(self) -> Path:
        path = Path(QStandardPaths.writableLocation(QStandardPaths.AppConfigLocation))
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _user_dictionary_path(self) -> Path:
        return self._config_dir() / "chemical_dictionary.json"

    def _legacy_user_dictionary_path(self) -> Path:
        return self._config_dir() / "classes.json"

    def _load_user_dictionary(self) -> list[dict[str, str]]:
        new_path = self._user_dictionary_path()
        legacy_path = self._legacy_user_dictionary_path()

        source = new_path if new_path.exists() else legacy_path if legacy_path.exists() else None
        if source is None:
            return [dict(row) for row in self.default_dictionary]

        try:
            with source.open("r", encoding="utf-8") as f:
                rows = json.load(f)
            rows = core.validate_dictionary(rows)
            # If we loaded the old classes.json, transparently migrate it to v2.
            if source == legacy_path:
                with new_path.open("w", encoding="utf-8") as f:
                    json.dump(rows, f, ensure_ascii=False, indent=2)
            return rows
        except Exception:
            return [dict(row) for row in self.default_dictionary]

    def _save_user_dictionary(self) -> None:
        rows = self._dictionary_from_table()
        rows = core.validate_dictionary(rows)
        with self._user_dictionary_path().open("w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2)
        self.dictionary_rows = rows

    # -------------------- interface --------------------

    def _build_calculation_tab(self) -> QWidget:
        root = QWidget()
        layout = QVBoxLayout(root)

        title = QLabel("RRF Calculator")
        title.setStyleSheet("font-size: 24px; font-weight: 600;")
        layout.addWidget(title)

        subtitle = QLabel(
            "Load a chromatographic XLSX file. The app resolves molecular formulas from the "
            "Chemical dictionary, calculates C, H, Molecular Weight, RRF, RRF*V and w, %, "
            "then creates the Class summary heatmap."
        )
        subtitle.setWordWrap(True)
        layout.addWidget(subtitle)

        file_row = QHBoxLayout()
        self.file_label = QLabel("No file selected")
        self.file_label.setTextInteractionFlags(Qt.TextSelectableByMouse)
        choose = QPushButton("Choose XLSX…")
        choose.clicked.connect(self._choose_file)
        file_row.addWidget(self.file_label, 1)
        file_row.addWidget(choose)
        layout.addLayout(file_row)

        button_row = QHBoxLayout()
        validate_btn = QPushButton("Validate")
        validate_btn.clicked.connect(self._validate_clicked)
        calculate_btn = QPushButton("Calculate & Save")
        calculate_btn.clicked.connect(self._calculate_clicked)
        button_row.addWidget(validate_btn)
        button_row.addWidget(calculate_btn)
        button_row.addStretch(1)
        layout.addLayout(button_row)

        self.status_label = QLabel("Status: waiting for a file")
        self.status_label.setStyleSheet("font-weight: 600;")
        layout.addWidget(self.status_label)

        self.log_box = QPlainTextEdit()
        self.log_box.setReadOnly(True)
        layout.addWidget(self.log_box, 1)
        return root

    def _build_dictionary_tab(self) -> QWidget:
        root = QWidget()
        layout = QVBoxLayout(root)

        info = QLabel(
            "Each row is either a class or an exact compound. "
            "Class example: aromatic | class | C[n]H[2*n-6]. "
            "Exact compound example: benzoic acid | compound | C7H6O2. "
            "Molecular weight is calculated automatically from the resolved formula."
        )
        info.setWordWrap(True)
        layout.addWidget(info)

        self.dict_table = QTableWidget(0, 4)
        self.dict_table.setHorizontalHeaderLabels(
            ["name", "type", "formula_rule", "summary_group"]
        )
        self.dict_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.dict_table, 1)

        controls = QHBoxLayout()
        add_btn = QPushButton("+ Add entry")
        add_btn.clicked.connect(self._add_dictionary_row)
        delete_btn = QPushButton("Delete selected")
        delete_btn.clicked.connect(self._delete_dictionary_rows)
        reset_btn = QPushButton("Restore defaults")
        reset_btn.clicked.connect(self._restore_defaults)
        save_btn = QPushButton("Save dictionary")
        save_btn.clicked.connect(self._save_dictionary_clicked)

        controls.addWidget(add_btn)
        controls.addWidget(delete_btn)
        controls.addStretch(1)
        controls.addWidget(reset_btn)
        controls.addWidget(save_btn)
        layout.addLayout(controls)
        return root

    # -------------------- drag and drop --------------------

    def dragEnterEvent(self, event):
        urls = event.mimeData().urls()
        if urls and urls[0].toLocalFile().lower().endswith(".xlsx"):
            event.acceptProposedAction()

    def dropEvent(self, event):
        self._set_input_file(Path(event.mimeData().urls()[0].toLocalFile()))

    # -------------------- dictionary --------------------

    def _make_type_combo(self, value: str = "class") -> QComboBox:
        combo = QComboBox()
        combo.addItems(["class", "compound"])
        index = combo.findText(value)
        combo.setCurrentIndex(index if index >= 0 else 0)
        return combo

    def _fill_dictionary_table(self) -> None:
        self.dict_table.setRowCount(0)
        for row in self.dictionary_rows:
            r = self.dict_table.rowCount()
            self.dict_table.insertRow(r)
            self.dict_table.setItem(r, 0, QTableWidgetItem(str(row["name"])))
            self.dict_table.setCellWidget(r, 1, self._make_type_combo(str(row["type"])))
            self.dict_table.setItem(r, 2, QTableWidgetItem(str(row["formula_rule"])))
            self.dict_table.setItem(r, 3, QTableWidgetItem(str(row["summary_group"])))
        self.dict_table.resizeColumnsToContents()

    def _dictionary_from_table(self) -> list[dict[str, str]]:
        rows = []
        for r in range(self.dict_table.rowCount()):
            name_item = self.dict_table.item(r, 0)
            formula_item = self.dict_table.item(r, 2)
            group_item = self.dict_table.item(r, 3)
            type_widget = self.dict_table.cellWidget(r, 1)
            entry_type = type_widget.currentText().strip() if isinstance(type_widget, QComboBox) else ""
            rows.append(
                {
                    "name": "" if name_item is None else name_item.text().strip(),
                    "type": entry_type,
                    "formula_rule": "" if formula_item is None else formula_item.text().strip(),
                    "summary_group": "" if group_item is None else group_item.text().strip(),
                }
            )
        return rows

    def _add_dictionary_row(self) -> None:
        r = self.dict_table.rowCount()
        self.dict_table.insertRow(r)
        self.dict_table.setItem(r, 0, QTableWidgetItem(""))
        self.dict_table.setCellWidget(r, 1, self._make_type_combo("class"))
        self.dict_table.setItem(r, 2, QTableWidgetItem(""))
        self.dict_table.setItem(r, 3, QTableWidgetItem(""))
        self.dict_table.setCurrentCell(r, 0)
        self.dict_table.editItem(self.dict_table.item(r, 0))

    def _delete_dictionary_rows(self) -> None:
        rows = sorted({index.row() for index in self.dict_table.selectedIndexes()}, reverse=True)
        for row in rows:
            self.dict_table.removeRow(row)

    def _save_dictionary_clicked(self) -> None:
        try:
            self._save_user_dictionary()
        except Exception as exc:
            QMessageBox.critical(self, "Dictionary error", str(exc))
            return
        self._log("Chemical dictionary saved.")
        QMessageBox.information(self, "Dictionary", "Chemical dictionary saved.")

    def _restore_defaults(self) -> None:
        answer = QMessageBox.question(
            self,
            "Restore defaults",
            "Replace the current Chemical dictionary with the built-in defaults?",
        )
        if answer != QMessageBox.Yes:
            return
        self.dictionary_rows = [dict(row) for row in self.default_dictionary]
        self._fill_dictionary_table()
        self._save_user_dictionary()
        self._log("Default Chemical dictionary restored.")

    # -------------------- file / validation --------------------

    def _choose_file(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(
            self,
            "Choose XLSX file",
            "",
            "Excel workbook (*.xlsx)",
        )
        if filename:
            self._set_input_file(Path(filename))

    def _set_input_file(self, path: Path) -> None:
        self.input_path = path
        self.file_label.setText(str(path))
        self.status_label.setText("Status: file selected")
        self._log(f"Selected: {path}")

    def _scan_file(self):
        if self.input_path is None:
            raise ValueError("Choose an XLSX file first.")

        dictionary_rows = core.validate_dictionary(self._dictionary_from_table())

        wb = load_workbook(self.input_path, read_only=True, data_only=False)
        ws = wb.active
        header_row, headers = core.find_header_row_and_columns(ws)
        group_col = headers["group name"]
        volume_col = headers["volume"]

        errors = []
        warnings = []
        unknown: dict[str, list[int]] = {}
        total_rows = 0
        no_c_rows = []
        suffix_rows = []

        for row_idx in range(header_row + 1, ws.max_row + 1):
            group_name = ws.cell(row_idx, group_col).value
            volume = ws.cell(row_idx, volume_col).value

            if group_name is None and volume is None:
                continue
            total_rows += 1

            if group_name is None or core.normalize_text(group_name) == "":
                errors.append(f"row {row_idx}: empty Group Name")
                continue

            try:
                float(volume)
            except (TypeError, ValueError):
                errors.append(f"row {row_idx}: invalid Volume = {volume!r}")
                continue

            try:
                resolved = core.resolve_group_name(group_name, dictionary_rows)
                if resolved.suffix:
                    suffix_rows.append((row_idx, core.normalize_text(group_name)))
            except core.MissingCarbonNumber:
                no_c_rows.append((row_idx, core.normalize_text(group_name)))
            except core.UnknownChemical as exc:
                text = str(exc)
                unknown.setdefault(text, []).append(row_idx)
            except Exception as exc:
                errors.append(f"row {row_idx}: {exc}")

        wb.close()

        for message, row_numbers in sorted(unknown.items()):
            preview = ", ".join(map(str, row_numbers[:8]))
            if len(row_numbers) > 8:
                preview += ", ..."
            errors.append(f"{message} Rows: {preview}.")

        if no_c_rows:
            warnings.append(
                f"{len(no_c_rows)} class row(s) have no explicit C number and will be skipped."
            )
        if suffix_rows:
            warnings.append(
                f"{len(suffix_rows)} row(s) contain '+' or '?'; numeric C will be used."
            )

        return {
            "ok": not errors,
            "rows": total_rows,
            "errors": errors,
            "warnings": warnings,
            "unknown": unknown,
        }

    def _show_scan(self, result) -> None:
        self._log("")
        self._log(f"Rows detected: {result['rows']}")
        self._log(f"Errors: {len(result['errors'])}")
        self._log(f"Warnings: {len(result['warnings'])}")
        for msg in result["errors"]:
            self._log("ERROR: " + msg)
        for msg in result["warnings"]:
            self._log("WARNING: " + msg)

        if result["ok"]:
            self.status_label.setText("Status: ready to calculate")
        else:
            self.status_label.setText("Status: action required")
            if result["unknown"]:
                self.tabs.setCurrentIndex(1)

    def _validate_clicked(self):
        try:
            result = self._scan_file()
        except Exception as exc:
            self.status_label.setText("Status: validation failed")
            self._log(f"ERROR: {exc}")
            QMessageBox.critical(self, "Validation error", str(exc))
            return None
        self._show_scan(result)
        return result

    # -------------------- calculation --------------------

    def _calculate_clicked(self) -> None:
        result = self._validate_clicked()
        if result is None or not result["ok"]:
            QMessageBox.warning(
                self,
                "Cannot calculate",
                "Fix validation errors first. Unknown entries can be added in Chemical dictionary.",
            )
            return

        default_output = self.input_path.with_name(self.input_path.stem + "_calculated.xlsx")
        filename, _ = QFileDialog.getSaveFileName(
            self,
            "Save calculated workbook",
            str(default_output),
            "Excel workbook (*.xlsx)",
        )
        if not filename:
            return

        output_path = Path(filename)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")

        try:
            self._save_user_dictionary()
            stats = core.calculate_file(self.input_path, self.dictionary_rows, output_path)
        except PermissionError:
            QMessageBox.critical(
                self,
                "Cannot save file",
                "The output workbook may be open in Excel. Close it and try again.",
            )
            return
        except Exception as exc:
            self._log(f"ERROR: {exc}")
            QMessageBox.critical(self, "Calculation error", str(exc))
            return

        self._log(f"Done: {output_path}")
        self._log(f"Calculated rows: {stats['calculated_rows']}")
        self._log(f"Sum(RRF*V): {stats['sum_rrf_v']:.6f}")
        self._log(f"Sum(w, %): {stats['sum_w']:.10f}")
        self._log(f"Summary sheet: {stats['summary_sheet']}")
        if stats["suffix_rows"]:
            self._log(
                f"WARNING: {len(stats['suffix_rows'])} row(s) contain '+' or '?'; numeric C was used."
            )
        if stats["skipped_no_c"]:
            self._log(
                f"WARNING: skipped {len(stats['skipped_no_c'])} class row(s) without explicit C number."
            )

        self.status_label.setText("Status: finished")
        QMessageBox.information(
            self,
            "Done",
            f"Calculation completed successfully.\n\n{output_path}",
        )

    def _log(self, text: str) -> None:
        self.log_box.appendPlainText(text)


def main() -> None:
    app = QApplication(sys.argv)
    app.setApplicationName("RRF Calculator")
    window = RRFWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
